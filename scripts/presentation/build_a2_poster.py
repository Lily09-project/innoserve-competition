"""Build an editable A2 methods poster from validated recorded aggregates."""
from pathlib import Path
import json
import sys

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Mm, Pt
from pptx.oxml.xmlchemy import OxmlElement

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'scripts'))
from validate_results import read_results, validate_results
from media_results import load_metrics
from finalize_delivery import normalize
from style_contract import font_runs

PALETTE = json.loads((ROOT/'design-system/palette.json').read_text(encoding='utf-8'))
BLUE, NAVY, GOLD, TEAL = (PALETTE[key].lstrip('#') for key in ('blue', 'navy', 'gold', 'teal'))
INK, MUTED, PALE, LINE = (PALETTE[key].lstrip('#') for key in ('ink', 'muted', 'gray', 'line'))
BLUE_TINT, TEAL_TINT, GOLD_TINT, BLUE_MUTED = (PALETTE[key].lstrip('#') for key in ('blue_tint', 'teal_tint', 'gold_tint', 'blue_muted'))
WHITE = 'FFFFFF'


def build():
    data = validate_results(read_results())
    m = load_metrics(ROOT)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Mm(420), Mm(594)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor.from_string(WHITE)
    for key in ('author', 'last_modified_by'):
        setattr(prs.core_properties, key, '')
    prs.core_properties.title = 'AI資安框架｜A2方法海報'

    def shape(kind, x, y, w, h, fill, stroke=None):
        s = slide.shapes.add_shape(kind, Mm(x), Mm(y), Mm(w), Mm(h))
        if fill:
            s.fill.solid()
            s.fill.fore_color.rgb = RGBColor.from_string(fill)
        else:
            s.fill.background()
        if stroke:
            s.line.color.rgb = RGBColor.from_string(stroke)
            s.line.width = Pt(.8)
        else:
            s.line.fill.background()
        s._element.spPr.append(OxmlElement('a:effectLst'))
        return s

    def rect(x, y, w, h, fill=PALE, stroke=None):
        return shape(MSO_SHAPE.RECTANGLE, x, y, w, h, fill, stroke)

    def card(x, y, w, h, fill=PALE, stroke=None):
        s = shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h, fill, stroke)
        s.adjustments[0] = min(.5, 2.4/min(w, h))
        return s

    def disc(x, y, d, fill, stroke=None):
        return shape(MSO_SHAPE.OVAL, x, y, d, d, fill, stroke)

    def text(x, y, w, h, content, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT):
        box = slide.shapes.add_textbox(Mm(x), Mm(y), Mm(w), Mm(h))
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = False
        tf.margin_left = tf.margin_right = 0
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        for i, paragraph in enumerate(content.split('\n')):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = align
            p.space_before = p.space_after = Pt(0)
            p.line_spacing = 1.1
            for token, family in font_runs(paragraph):
                r = p.add_run()
                r.text = token
                r.font.name = family
                r.font.size, r.font.bold = Pt(size), bold
                r.font.color.rgb = RGBColor.from_string(color)
        return box

    def line(x1, y1, x2, y2, color=LINE, width=1):
        s = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Mm(x1), Mm(y1), Mm(x2), Mm(y2))
        s.line.color.rgb = RGBColor.from_string(color)
        s.line.width = Pt(width)
        s._element.spPr.append(OxmlElement('a:effectLst'))

    def arrow(x, y, w, h, kind=MSO_SHAPE.RIGHT_ARROW, color=BLUE):
        return shape(kind, x, y, w, h, color)

    def icon(n, x, y, size=23):
        slide.shapes.add_picture(str(ROOT / f'deliverables/source/flat-icon-{n}.png'), Mm(x), Mm(y), width=Mm(size), height=Mm(size))

    def document(x, y, w, h):
        rect(x, y, w, h, WHITE, BLUE)
        for k in range(3):
            yy = y+h*(k+1)/4
            line(x+w*.22, yy, x+w*.78, yy, BLUE, .8)

    def section(y, number, title):
        text(24, y-1.5, 16, 13, number, 22, TEAL, True)
        text(41, y-1.5, 300, 13, title, 26, NAVY, True)
        line(24, y+13, 396, y+13)

    # Light masthead matching the deck cover: formal name, one takeaway and the four-stage evidence chain.
    rect(0, 0, 420, 162, PALE)
    line(0, 162, 420, 162, LINE, 1.5)
    text(24, 12, 300, 8, '2026 資安技術組 · 決賽作品', 18, TEAL, True)
    title = text(24, 20, 372, 71, '融合新型內容相依時序方向模型\n入侵偵測與任務導向檢索增強稽核\n之 AI 資安框架', 46, NAVY, True)
    title.name = 'PosterTitle_Masthead'
    rect(24, 94, 24, 1.6, GOLD)
    text(24, 97, 372, 11, '從流量告警，走到有證據、可回復的處置', 27, BLUE, True)
    text(24, 109, 372, 8, '同一事件串起模型判定、相似案例、AI稽核與處置回條。', 18, MUTED)
    line(56, 136, 364, 136, LINE, 2)
    for i, (label, image) in enumerate((('偵測', 2), ('檢索', 3), ('稽核', None), ('回復', 4))):
        cx = 56+i*102.7
        disc(cx-8, 128, 16, WHITE, BLUE)
        if image:
            icon(image, cx-5, 131, 10)
        else:
            document(cx-3.2, 131, 6.4, 10)
        text(cx-20, 146, 40, 9, label, 20, NAVY, True, PP_ALIGN.CENTER)

    section(169, '01', '四階段方法，串成同一條事件鏈')
    stages = [('互補判定', '內容＋時序', 2), ('任務檢索', '表徵＋機率', 3),
              ('證據稽核', '保留正式分類', None), ('受控回復', '策略＋期限', 4)]
    stage_outputs = ['產生正式類別、\n信心與流量表徵',
                     '依事件任務找案例\n附上相似類別證據',
                     '核對判定與案例\n附加三態稽核結論',
                     '策略允許才介入\n讀回並確認回復']
    for i, (name, detail, image) in enumerate(stages):
        y = 187+i*47
        card(24, y, 372, 43, PALE)
        if image:
            icon(image, 29, y+5, 13)
        else:
            document(31, y+6, 9, 12)
        text(46, y+4, 72, 9, f'{i+1:02d}  '+name, 22, NAVY, True)
        text(46, y+13, 72, 6, detail, 14, TEAL, True)
        text(30, y+23, 90, 15, stage_outputs[i], 15, INK)
        if i < 3:
            arrow(258, y+43.3, 7, 3.4, MSO_SHAPE.DOWN_ARROW, BLUE_MUTED)

    def node(x, y, w, h, label, sub=None, fill=WHITE, color=BLUE):
        card(x, y, w, h, fill, LINE if fill == WHITE else None)
        text(x, y+(1 if sub else 0), w, h*(.56 if sub else 1), label, 19, color, True, PP_ALIGN.CENTER)
        if sub:
            text(x, y+h*.54, w, h*.43, sub, 13.5, MUTED, align=PP_ALIGN.CENTER)

    def merge(x, y1, y2, ym):
        line(x, y1, x+8, y1, BLUE, 1.2)
        line(x, y2, x+8, y2, BLUE, 1.2)
        line(x+8, y1, x+8, y2, BLUE, 1.2)
        arrow(x+9, ym-2.5, 7, 5)

    # 01: parallel models enter one calibrated output.
    node(126, 191, 78, 16, 'CTBM', '內容·時間差·方向')
    node(126, 211, 78, 16, 'PBCNN', '短窗位元組模式')
    merge(204, 199, 219, 209)
    node(222, 196, 84, 26, '校準＋融合', 'OOF Gate', BLUE_TINT)
    arrow(308, 206.5, 10, 5)
    node(320, 196, 72, 26, '17類判定', '信心與512維表徵', TEAL_TINT, TEAL)
    # 02: task vector composition and train-only evidence retrieval.
    node(126, 238, 78, 16, '512維表徵')
    node(126, 258, 78, 16, '17維機率')
    merge(204, 246, 266, 256)
    node(222, 243, 84, 26, '529維向量', 'L2正規化', BLUE_TINT)
    arrow(308, 253.5, 10, 5)
    card(320, 243, 72, 26, TEAL_TINT)
    text(320, 244, 72, 10, 'Top-5訓練案例', 18, TEAL, True, PP_ALIGN.CENTER)
    for rank in range(5):
        x = 323+rank*13.4
        card(x, 255, 11, 10, WHITE)
        text(x, 255, 11, 10, str(rank+1), 16, TEAL, True, PP_ALIGN.CENTER)
    # 03: attach an audit; never overwrite the original classifier.
    node(126, 285, 78, 16, '正式判定')
    node(126, 305, 78, 16, '案例證據')
    merge(204, 293, 313, 303)
    node(222, 290, 84, 26, '本機AI稽核', '核對同一事件', BLUE_TINT)
    arrow(308, 300.5, 10, 5)
    for k, label in enumerate(('支持', '複核', '不確定')):
        node(320, 288+k*10.3, 72, 9, label, None, TEAL_TINT, TEAL)
    # 04: allowed action, readback and expiry, with the review branch explicit.
    text(126, 330, 266, 8, '策略檢查：範圍、動作與期限｜來源限速等', 16, INK)
    for k, label in enumerate(('套用規則', '讀回驗證', '到期回復')):
        x = 126+k*96
        node(x, 340, 78, 14, label, None, TEAL_TINT if k == 2 else WHITE, TEAL if k == 2 else BLUE)
        if k < 2:
            arrow(x+81, 344.5, 12, 5)
    text(126, 356, 266, 7, '複核／不確定 → 保留監測或蒐證', 14, MUTED)
    card(24, 377, 372, 13, NAVY)
    for i in range(4):
        disc(70+i*93, 375, 4, GOLD)
    text(24, 377, 372, 13, '同一事件識別碼：正式分類／案例證據／稽核結論／處置回條', 18, WHITE, True, PP_ALIGN.CENTER)

    section(398, '02', '三項優勢，從判定走到可追溯處置')
    innovations = [('互補判定', '內容、時間差與方向互補\n校準融合，固定17類輸出\nMacro-F1達93.05%'),
                   ('任務證據', '529維任務向量找Top-5案例\nAI核對判定，保留正式分類\n命中率97.62%'),
                   ('受控可追溯', '本機AI，事件資料留在本機\n同事件串接判定、證據與回條\n策略檢查、讀回及到期回復')]
    for i, (heading, copy) in enumerate(innovations):
        x = 24+i*126
        card(x, 416, 120, 44, WHITE, LINE)
        rect(x+6, 420, 14, 1.4, GOLD)
        text(x+6, 423, 108, 10, heading, 22, TEAL, True)
        text(x+6, 434, 110, 23, copy, 15, INK)

    section(468, '03', '以分階段實測呈現效果')
    text(24, 485, 184, 7, f'相同{m["detection_n"]:,}筆測試的 Macro-F1', 15, MUTED)
    for i, name in enumerate(('PBCNN', 'CTBM', 'OOF Gate')):
        y = 494+i*12
        value = data['detection'][name]['macro_f1']*100
        best = i == 2
        text(24, y, 40, 8, name, 17, TEAL if best else INK, True)
        rect(66, y+1, 94, 6, PALE)
        rect(66, y+1, 94*value/100, 6, TEAL if best else BLUE_MUTED)
        text(163, y, 43, 8, f'{value:.2f}%', 20, TEAL if best else INK, True, PP_ALIGN.RIGHT)
    for tick in (0, 50, 100):
        x = 66 + 94*tick/100
        line(x, 530, x, 532, LINE, .6)
        text(x-9, 532, 18, 5, str(tick)+'%', 12, MUTED, align=PP_ALIGN.CENTER)
    text(24, 539, 184, 7, f'整體準確率{m["accuracy"]*100:.2f}%·固定17類分類', 15, MUTED)
    line(213, 485, 213, 546)
    r, j = data['retrieval'], data['judge']
    text(220, 485, 176, 7, f'{m["bank_rows"]:,}筆訓練案例／{m["queries"]}筆查詢', 15, MUTED)
    text(220, 492, 176, 14, f'Top-5  {r["top5"]*100:.2f}%', 34, TEAL, True)
    text(220, 506, 176, 7, f'MRR@5  {r["mrr5"]*100:.2f}%·訓練／查詢重疊{m["overlap"]}', 14, MUTED)
    for i, (name, key) in enumerate((('支持', 'support'), ('複核', 'review'), ('不確定', 'uncertain'))):
        x = 220+i*59.3
        card(x, 515, 55, 22, TEAL_TINT if key == 'support' else PALE)
        text(x, 515.5, 55, 12, str(j['assessment_counts'][key]), 26, TEAL if key == 'support' else NAVY, True, PP_ALIGN.CENTER)
        text(x, 528, 55, 8, name, 15, INK, align=PP_ALIGN.CENTER)
    text(220, 539, 176, 7, f'{m["errors"]}筆偵測誤判中{m["intercepted"]}筆分流；其餘{m["false_support"]}筆獲支持', 13.5, MUTED)
    card(24, 550, 372, 22, PALE)
    text(28, 551, 64, 20, f'{m["live_classes"]}類本機重跑\n{m["network_cases"]}種隔離TCP情境', 15, NAVY, True)
    case_labels = {'BENIGN': '正常', 'FTP-BruteForce': 'FTP', 'SSH-Bruteforce': 'SSH',
                   'DDoS-LOIC-HTTP': 'HTTP', 'Web-XSS': 'XSS', 'Bot': 'Bot'}
    for i, case in enumerate(data['framework']['cases']):
        x = 94+i*33
        active = case['rules_applied'] > 0
        card(x, 553, 30, 16, GOLD_TINT if active else WHITE, None if active else LINE)
        text(x, 553.5, 30, 8, case_labels[case['validation_label']], 15, NAVY, True, PP_ALIGN.CENTER)
        status = '規則' if active else '監測' if case['action'] == 'monitor' else '蒐證'
        text(x, 561, 30, 7, status, 13, MUTED, align=PP_ALIGN.CENTER)
    framework = data['framework']
    text(296, 551, 96, 20, f'背景保留 {framework["background_preserved_cases"]}/{m["network_cases"]}\n到期回復 {framework["rollback_passed_cases"]}/{m["network_cases"]}', 18, TEAL, True, PP_ALIGN.RIGHT)
    text(24, 576, 372, 7, f'{m["rules_cases"]}種套用規則、{m["passive_cases"]}種監測或蒐證；{m["rollback"]}/{m["network_cases"]}為隔離交易與回復結果，非原生攻擊防禦率。', 15, MUTED, align=PP_ALIGN.CENTER)

    target = ROOT / 'deliverables/決賽海報_AI資安框架_A2.pptx'
    intermediate = ROOT / 'validation/poster-a2/authoring.pptx'
    intermediate.parent.mkdir(parents=True, exist_ok=True)
    prs.save(intermediate)
    normalize(intermediate, target)
    print(target.name)


if __name__ == '__main__':
    build()
